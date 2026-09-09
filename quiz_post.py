#!/usr/bin/env python3
"""dlab記事ベースの「今日のクイズ」をSlackに投稿する。

data/quiz/quiz_bank.xlsx の quiz_bank シートから1記事分(4問)、
psychology_bonus シートから心理学ボーナス1問を、同じ day_index で取り出し、
1) 問題本文(4問+心理学ボーナス1問)を投稿
2) そのスレッドに正解・解説を返信
する。全記事を投稿し終えたら quiz_posted.log をリセットして最初から繰り返す。
"""
import json
import os
import urllib.request

import openpyxl

QUIZ_BANK_PATH = "data/quiz/quiz_bank.xlsx"
POSTED_LOG_PATH = "quiz_posted.log"

SLACK_API_URL = "https://slack.com/api/chat.postMessage"


def load_quiz_bank():
    wb = openpyxl.load_workbook(QUIZ_BANK_PATH, read_only=True)

    ws = wb["quiz_bank"]
    articles = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        (day_index, channel, channel_label, title, url, published, q_no,
         question, choice_a, choice_b, choice_c, choice_d, correct,
         explanation) = row
        article = articles.setdefault(day_index, {
            "day_index": day_index,
            "channel_label": channel_label,
            "title": title,
            "url": url,
            "published": published,
            "questions": [],
        })
        article["questions"].append({
            "q_no": q_no,
            "question": question,
            "choices": {"A": choice_a, "B": choice_b, "C": choice_c, "D": choice_d},
            "correct": correct,
            "explanation": explanation,
        })

    ws2 = wb["psychology_bonus"]
    bonus_by_day = {}
    for row in ws2.iter_rows(min_row=2, values_only=True):
        (day_index, title, url, published, question, choice_a, choice_b,
         choice_c, choice_d, correct, explanation) = row
        bonus_by_day[day_index] = {
            "title": title,
            "url": url,
            "published": published,
            "question": question,
            "choices": {"A": choice_a, "B": choice_b, "C": choice_c, "D": choice_d},
            "correct": correct,
            "explanation": explanation,
        }

    for day_index, article in articles.items():
        article["psychology_bonus"] = bonus_by_day.get(day_index)

    return [articles[k] for k in sorted(articles)]


def load_posted_days():
    if not os.path.exists(POSTED_LOG_PATH):
        return set()
    with open(POSTED_LOG_PATH, encoding="utf-8") as f:
        return {int(line.strip()) for line in f if line.strip()}


def pick_next_article(articles, posted_days):
    for article in articles:
        if article["day_index"] not in posted_days:
            return article
    return None


def build_quiz_message(article):
    lines = [f"🧠 *今日のクイズ*（テーマ: {article['channel_label']}）", ""]
    lines.append(f"元記事: <{article['url']}|{article['title']}>（{article['published']}時点の情報）")
    lines.append("")
    for q in article["questions"]:
        lines.append(f"*Q{q['q_no']}.* {q['question']}")
        lines.append(
            f"A) {q['choices']['A']}　B) {q['choices']['B']}　"
            f"C) {q['choices']['C']}　D) {q['choices']['D']}"
        )
        lines.append("")

    bonus = article.get("psychology_bonus")
    if bonus:
        lines.append(f"🧩 *心理学ボーナス問題*")
        lines.append(f"出典: <{bonus['url']}|{bonus['title']}>（{bonus['published']}時点の情報）")
        lines.append(f"*Q5.* {bonus['question']}")
        lines.append(
            f"A) {bonus['choices']['A']}　B) {bonus['choices']['B']}　"
            f"C) {bonus['choices']['C']}　D) {bonus['choices']['D']}"
        )
        lines.append("")

    lines.append("正解はこのスレッドの返信をチェック👇")
    return "\n".join(lines)


def build_answer_message(article):
    lines = ["✅ *正解発表*", ""]
    for q in article["questions"]:
        lines.append(f"Q{q['q_no']}: *{q['correct']}* ({q['choices'][q['correct']]})")
        lines.append(f"　→ {q['explanation']}")

    bonus = article.get("psychology_bonus")
    if bonus:
        lines.append(f"Q5(心理学ボーナス): *{bonus['correct']}* ({bonus['choices'][bonus['correct']]})")
        lines.append(f"　→ {bonus['explanation']}")

    return "\n".join(lines)


def slack_post(token, channel, text, thread_ts=None):
    payload = {"channel": channel, "text": text}
    if thread_ts:
        payload["thread_ts"] = thread_ts
    req = urllib.request.Request(
        SLACK_API_URL,
        data=json.dumps(payload).encode("utf-8"),
        headers={
            "Authorization": f"Bearer {token}",
            "Content-Type": "application/json; charset=utf-8",
        },
        method="POST",
    )
    with urllib.request.urlopen(req) as resp:
        result = json.loads(resp.read().decode("utf-8"))
    if not result.get("ok"):
        raise RuntimeError(f"Slack API error: {result}")
    return result


def main():
    token = os.environ["SLACK_BOT_TOKEN_2"]
    channel = os.environ["SLACK_CHANNEL_QUIZ"]

    articles = load_quiz_bank()
    posted_days = load_posted_days()
    article = pick_next_article(articles, posted_days)

    if article is None:
        # 全部投稿し終えたのでリセットして最初の記事から再開する
        posted_days = set()
        article = articles[0]

    quiz_text = build_quiz_message(article)
    result = slack_post(token, channel, quiz_text)
    thread_ts = result["ts"]

    answer_text = build_answer_message(article)
    slack_post(token, channel, answer_text, thread_ts=thread_ts)

    posted_days.add(article["day_index"])
    with open(POSTED_LOG_PATH, "w", encoding="utf-8") as f:
        for day in sorted(posted_days):
            f.write(f"{day}\n")

    print(f"posted day_index={article['day_index']} title={article['title']}")


if __name__ == "__main__":
    main()
